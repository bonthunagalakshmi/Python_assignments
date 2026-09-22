"""Behavioral checks for DriveEase.

The suite uses a temporary workbook and never touches the user's real data file.
Run with: python test_system.py
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import load_workbook

from data.excel_manager import ExcelManager
from exceptions.rental_exceptions import PaymentError, VehicleUnavailableError
from services.customer_service import CustomerService
from services.fleet_service import FleetService
from services.maintenance_service import MaintenanceService
from services.payment_service import PaymentService
from services.pricing_service import PricingService
from services.rental_service import RentalService
from services.report_service import ReportService


checks = 0


def check(condition: bool, message: str) -> None:
    global checks
    assert condition, message
    checks += 1


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="driveease-test-") as folder:
        workbook_path = Path(folder) / "test.xlsx"
        excel = ExcelManager(workbook_path, seed_sample=False)
        check(workbook_path.exists(), "workbook was created")
        workbook = load_workbook(workbook_path, read_only=True)
        check({"Customers", "Vehicles", "Rentals", "Payments", "Invoices", "Maintenance"} <=
              set(workbook.sheetnames), "all workbook sheets are created")
        workbook.close()

        customers = CustomerService(excel)
        fleet = FleetService(excel)
        pricing = PricingService()
        payments = PaymentService(excel)
        rentals = RentalService(excel, customers, fleet, pricing, payments)
        reports = ReportService(excel)
        maintenance = MaintenanceService(excel, fleet)

        customer = customers.create_customer(
            "Test Driver", "driver@example.com", "+91 90000 00000",
            "DL-TEST-1", "1 Test Street",
        )
        customer_two = customers.create_customer(
            "Second Driver", "second@example.com", "+91 90000 00001",
            "DL-TEST-2", "2 Test Street",
        )
        check(customer.customer_id == "CUS1001", "customer IDs start at CUS1001")
        check(customer_two.customer_id == "CUS1002", "customer IDs increment without duplication")
        check(len(customers.search_customers("test driver")) == 1, "customer search is case-insensitive")

        car = fleet.add_vehicle("Car", "TEST-CAR-1", "Test", "Sedan", 2024, "Petrol", "Automatic", 5, 1000)
        bike = fleet.add_vehicle("Bike", "TEST-BIKE-1", "Test", "Bike", 2024, "Petrol", "Manual", 2, 500)
        van = fleet.add_vehicle("Van", "TEST-VAN-1", "Test", "Van", 2024, "Diesel", "Manual", 8, 2000)
        check(car.vehicle_id == "CAR1001" and bike.vehicle_id == "BIKE1001" and van.vehicle_id == "VAN1001",
              "vehicle IDs use type prefixes")
        check(len(fleet.search("test")) == 3, "vehicle search finds all test fleet records")
        check(round(car.calculate_rental_cost(3), 2) == 3000, "car polymorphic pricing")
        check(round(bike.calculate_rental_cost(3), 2) == 1350, "bike multi-day discount pricing")
        check(round(van.calculate_rental_cost(3), 2) == 6450, "van handling pricing")
        check(pricing.rental_days("2030-09-20", "2030-09-25") == 5, "rental days are exclusive of return date")

        available = rentals.available_vehicles("2030-09-20", "2030-09-25")
        check(len(available) == 3, "all non-maintenance vehicles are initially available")

        try:
            rentals.confirm_rental(
                customer.customer_id, car.vehicle_id, "2030-09-20", "2030-09-25",
                "UPI", "not-a-vpa",
            )
            raise AssertionError("invalid payment should fail")
        except PaymentError:
            pass
        check(len(rentals.list_rentals()) == 0, "failed payment creates no rental")
        check(len(payments.list_payments()) == 0, "failed payment creates no payment record")

        rental, invoice, payment = rentals.confirm_rental(
            customer.customer_id, car.vehicle_id, "2030-09-20", "2030-09-25",
            "CARD", "****4242", payment_amount=1000,
        )
        check(rental.status == "CONFIRMED", "successful payment confirms rental")
        check(invoice["payment_status"] == "PARTIAL", "partial initial payment updates invoice")
        check(payment.payment_method == "CARD", "card payment processor records method")
        check(not rentals.is_vehicle_available(car.vehicle_id, "2030-09-22", "2030-09-24"),
              "overlapping rental is rejected by date rule")
        try:
            rentals.create_quotation(
                customer_two.customer_id, car.vehicle_id, "2030-09-22", "2030-09-24"
            )
            raise AssertionError("overlapping quotation should fail")
        except VehicleUnavailableError:
            pass
        check(True, "overlap raises VehicleUnavailableError")
        check(rentals.is_vehicle_available(car.vehicle_id, "2030-09-25", "2030-09-28"),
              "end-inclusive boundary is available under start-inclusive/end-exclusive rule")
        check(fleet.get_vehicle(car.vehicle_id).status == "RESERVED", "future rental sets RESERVED status")

        _, boundary_invoice, boundary_payment = rentals.confirm_rental(
            customer_two.customer_id, car.vehicle_id, "2030-09-25", "2030-09-28",
            "UPI", "driver@upi",
        )
        check(boundary_invoice["payment_status"] == "PAID" and boundary_payment.payment_method == "UPI",
              "UPI payment and boundary booking succeed")

        maintenance_record = maintenance.start_maintenance(van.vehicle_id, "Brake inspection", 850, "Replace pads")
        check(fleet.get_vehicle(van.vehicle_id).status == "MAINTENANCE", "maintenance changes vehicle status")
        check(not rentals.is_vehicle_available(van.vehicle_id, "2030-09-20", "2030-09-25"),
              "maintenance vehicle cannot be rented")
        maintenance.complete_maintenance(maintenance_record["maintenance_id"])
        check(fleet.get_vehicle(van.vehicle_id).status == "AVAILABLE", "completed maintenance returns vehicle to service")

        return_rental, _, _ = rentals.confirm_rental(
            customer.customer_id, bike.vehicle_id, "2025-01-10", "2025-01-12",
            "CASH", payment_amount=500,
        )
        check(return_rental.status == "CONFIRMED", "cash payment supports partial booking payment")
        completed, returned_invoice, charges = rentals.return_vehicle(
            return_rental.rental_id, "2025-01-14", 1200, 40, "DAMAGED", "Scratch on door"
        )
        check(completed["status"] == "COMPLETED", "return completes rental")
        check(charges["late_fee"] > 0, "late return charge is calculated")
        check(charges["fuel_charge"] > 0, "fuel shortage charge is calculated")
        check(charges["damage_charge"] > 0, "damage charge is calculated")
        check(float(returned_invoice["amount_due"]) > 0 and returned_invoice["payment_status"] == "PARTIAL",
              "return charges create a partial invoice balance")
        payment_after_return, paid_invoice = rentals.collect_payment(
            return_rental.rental_id, float(returned_invoice["amount_due"]), "CASH", "CASH-FINAL"
        )
        check(payment_after_return.status == "SUCCESS" and paid_invoice["payment_status"] == "PAID",
              "final payment closes the invoice")
        check(fleet.get_vehicle(bike.vehicle_id).status == "AVAILABLE", "returned vehicle becomes available")
        check(rentals.is_vehicle_available(bike.vehicle_id, "2030-01-01", "2030-01-03"),
              "completed rental does not block future availability")

        reports_summary = reports.dashboard_summary()
        check(reports_summary["total_vehicles"] == 3, "reports count actual vehicles")
        check(reports_summary["revenue"] > 0, "reports calculate actual revenue")
        check(reports.most_rented_vehicles()[0]["rental_count"] >= 1, "most rented report uses rentals")
        check("CASH" in reports.revenue_by_payment_method(), "payment method report uses payment records")
        check(reports.customer_rental_history(customer.customer_id), "customer history report returns rentals")

        reopened = ExcelManager(workbook_path, seed_sample=False)
        check(len(reopened.read_records("Customers")) == 2, "customers persist across restart")
        check(len(reopened.read_records("Vehicles")) == 3, "vehicles persist across restart")
        check(len(reopened.read_records("Rentals")) == 3, "rentals persist across restart")
        check(len(reopened.read_records("Payments")) >= 4, "multiple payments persist across restart")
        check(len(reopened.read_records("Invoices")) == 3, "invoices persist across restart")
        check(len(reopened.read_records("Maintenance")) == 1, "maintenance persists across restart")

        sample_path = Path(folder) / "sample.xlsx"
        seeded = ExcelManager(sample_path, seed_sample=True)
        first_sample_counts = (len(seeded.read_records("Customers")), len(seeded.read_records("Vehicles")))
        reopened_seeded = ExcelManager(sample_path, seed_sample=True)
        check(first_sample_counts == (3, 7), "first launch seeds the documented sample dataset")
        check((len(reopened_seeded.read_records("Customers")), len(reopened_seeded.read_records("Vehicles"))) == first_sample_counts,
              "reopening does not duplicate sample data")

    print(f"{checks}/{checks} checks passed.")


if __name__ == "__main__":
    main()