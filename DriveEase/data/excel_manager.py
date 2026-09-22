from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook


SHEET_HEADERS: dict[str, list[str]] = {
    "Customers": [
        "customer_id", "name", "email", "phone", "driving_licence",
        "address", "registration_date", "status",
    ],
    "Vehicles": [
        "vehicle_id", "registration_number", "vehicle_type", "brand", "model",
        "year", "fuel_type", "transmission", "seats", "daily_rate",
        "odometer", "fuel_level", "condition", "status",
    ],
    "Rentals": [
        "rental_id", "customer_id", "vehicle_id", "pickup_date", "return_date",
        "actual_return_date", "base_amount", "additional_charges", "discount",
        "tax", "final_amount", "status", "created_date", "pickup_odometer",
        "pickup_fuel_level", "pickup_condition", "pickup_remarks",
        "return_odometer", "return_fuel_level", "return_condition",
        "return_remarks", "charge_breakdown",
    ],
    "Payments": [
        "payment_id", "rental_id", "amount", "payment_method", "payment_date",
        "status", "reference_id",
    ],
    "Invoices": [
        "invoice_id", "rental_id", "customer_id", "vehicle_id", "base_amount",
        "additional_charges", "discount", "tax", "total", "amount_paid",
        "amount_due", "payment_status", "generated_date",
    ],
    "Maintenance": [
        "maintenance_id", "vehicle_id", "date", "reason", "cost", "notes", "status",
    ],
}


class ExcelManager:
    """Small persistence gateway that preserves records across application restarts."""

    def __init__(self, path: str | Path = "DriveEase_Data.xlsx", seed_sample: bool = True):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        is_new = not self.path.exists()
        if is_new:
            self._create_workbook()
        else:
            self._repair_workbook()
        if is_new and seed_sample:
            self.seed_sample_data()

    def _create_workbook(self) -> None:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        for sheet_name, headers in SHEET_HEADERS.items():
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(headers)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = cell.font.copy(bold=True)
        workbook.save(self.path)

    def _repair_workbook(self) -> None:
        workbook = load_workbook(self.path)
        changed = False
        for sheet_name, headers in SHEET_HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                workbook.create_sheet(sheet_name).append(headers)
                changed = True
            else:
                sheet = workbook[sheet_name]
                if sheet.max_row == 1 and [cell.value for cell in sheet[1]] != headers:
                    sheet.delete_rows(1)
                    sheet.insert_rows(1)
                    for index, header in enumerate(headers, 1):
                        sheet.cell(1, index).value = header
                    changed = True
        if changed:
            workbook.save(self.path)

    def read_records(self, sheet_name: str) -> list[dict[str, object]]:
        if sheet_name not in SHEET_HEADERS:
            raise KeyError(f"Unknown sheet: {sheet_name}")
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        headers = [str(cell.value) for cell in sheet[1]]
        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None and str(value) != "" for value in row):
                continue
            record = {header: (row[index] if index < len(row) and row[index] is not None else "")
                      for index, header in enumerate(headers)}
            records.append(record)
        workbook.close()
        return records

    def append_record(self, sheet_name: str, record: dict[str, object]) -> None:
        workbook = load_workbook(self.path)
        sheet = workbook[sheet_name]
        headers = SHEET_HEADERS[sheet_name]
        sheet.append([record.get(header, "") for header in headers])
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(self.path)

    def replace_records(self, sheet_name: str, records: Iterable[dict[str, object]]) -> None:
        workbook = load_workbook(self.path)
        sheet = workbook[sheet_name]
        sheet.delete_rows(2, max(sheet.max_row - 1, 1))
        headers = SHEET_HEADERS[sheet_name]
        for record in records:
            sheet.append([record.get(header, "") for header in headers])
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(self.path)

    def update_record(self, sheet_name: str, key: str, value: object, updates: dict[str, object]) -> bool:
        workbook = load_workbook(self.path)
        sheet = workbook[sheet_name]
        headers = [str(cell.value) for cell in sheet[1]]
        if key not in headers:
            raise KeyError(f"Unknown column: {key}")
        key_index = headers.index(key) + 1
        changed = False
        for row in sheet.iter_rows(min_row=2):
            if row[key_index - 1].value == value:
                for field, field_value in updates.items():
                    if field in headers:
                        row[headers.index(field)].value = field_value
                changed = True
                break
        if changed:
            workbook.save(self.path)
        workbook.close()
        return changed

    def seed_sample_data(self) -> None:
        from datetime import date

        if self.read_records("Customers") or self.read_records("Vehicles"):
            return
        today = date.today().isoformat()
        customers = [
            {"customer_id": "CUS1001", "name": "Aarav Mehta", "email": "aarav@example.com",
             "phone": "+91 98765 12001", "driving_licence": "DL-MH-1001",
             "address": "12 Carter Road, Mumbai", "registration_date": today, "status": "ACTIVE"},
            {"customer_id": "CUS1002", "name": "Nisha Kapoor", "email": "nisha@example.com",
             "phone": "+91 98765 12002", "driving_licence": "DL-DL-1002",
             "address": "44 Hauz Khas, New Delhi", "registration_date": today, "status": "ACTIVE"},
            {"customer_id": "CUS1003", "name": "Rohan Iyer", "email": "rohan@example.com",
             "phone": "+91 98765 12003", "driving_licence": "DL-KA-1003",
             "address": "8 Indiranagar, Bengaluru", "registration_date": today, "status": "ACTIVE"},
        ]
        for record in customers:
            self.append_record("Customers", record)
        vehicles = [
            ("CAR1001", "MH01AB1234", "Car", "Maruti", "Baleno", 2023, "Petrol", "Automatic", 5, 2200, 18200, 82, "GOOD"),
            ("CAR1002", "DL04CD5678", "Car", "Hyundai", "Creta", 2022, "Diesel", "Automatic", 5, 3200, 26700, 76, "GOOD"),
            ("CAR1003", "KA05EF9012", "Car", "Honda", "City", 2021, "Petrol", "Manual", 5, 2500, 31400, 90, "GOOD"),
            ("BIKE1001", "MH12GH3456", "Bike", "Royal Enfield", "Classic 350", 2023, "Petrol", "Manual", 2, 900, 8400, 68, "GOOD"),
            ("BIKE1002", "KA03IJ7890", "Bike", "TVS", "Ntorq", 2024, "Petrol", "Automatic", 2, 650, 3200, 95, "GOOD"),
            ("VAN1001", "DL01KL2468", "Van", "Force", "Traveller", 2020, "Diesel", "Manual", 12, 5200, 48600, 72, "GOOD"),
            ("VAN1002", "MH14MN1357", "Van", "Tata", "Winger", 2022, "Diesel", "Manual", 9, 4400, 22100, 88, "GOOD"),
        ]
        for item in vehicles:
            fields = ["vehicle_id", "registration_number", "vehicle_type", "brand", "model", "year",
                      "fuel_type", "transmission", "seats", "daily_rate", "odometer", "fuel_level", "condition"]
            self.append_record("Vehicles", dict(zip(fields, item)) | {"status": "AVAILABLE"})